"""scripts/gen-service-objects.py — single source of truth for the Phase 2 service
objects (D-072, docs/service-data-model.md).

Run from the repo root:  python scripts/gen-service-objects.py   (needs openpyxl)
Emits, from ONE field spec, straight into the repo:
  salesforce/service-objects/{package.xml, objects/*.object, permissionsets/*.permissionset}
  sql/sundial_*_cache.sql          (7 Supabase cache tables, roofing-cache pattern)
  docs/*_Fields_by_Section.xlsx    (generator-format workbooks for the portal generator)
  salesforce/service-objects/spec.json (registry keys / tables / parent filters)
Change the model HERE (and in the data-model doc), regenerate, never hand-edit outputs.
Pass an output root as argv[1] to write somewhere else.
Pass --no-commercial to omit the three Sundial_Commercial__c lookups (use when the verify
script reports that object is absent from the org - a lookup to a missing object fails
the whole deploy). Re-run WITHOUT the flag once Sundial_Commercial__c exists and deploy
the three fields then.
"""
import html, os, json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NO_COMMERCIAL = "--no-commercial" in sys.argv
_args = [a for a in sys.argv[1:] if not a.startswith("--")]
ROOT = _args[0] if _args else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PKG = f"{ROOT}/salesforce/service-objects"
for d in (f"{PKG}/objects", f"{PKG}/permissionsets", f"{ROOT}/sql", f"{ROOT}/docs"):
    os.makedirs(d, exist_ok=True)

def esc(s): return html.escape(s, quote=False)

BILL_TO = ["Customer", "Internal Warranty", "Manufacturer", "Leasing Partner", "Other"]
TENANT = "Sundial_Tenant__c"

# ---------------------------------------------------------------------------
# Field spec helper. cache=True → column in the Supabase cache table.
# sys=True → system-maintained (Lambda/Flow-written), yellow in the workbook.
# ---------------------------------------------------------------------------
class F:
    def __init__(self, api, label, section, ftype, desc="", cache=True, sys=False, **kw):
        self.api, self.label, self.ftype, self.section, self.desc = api, label, ftype, section, desc
        self.cache, self.sys, self.kw = cache, sys, kw

def picklist_xml(values, default=None, restricted=True):
    vs = "\n".join(f"""                <value>
                    <fullName>{esc(v)}</fullName>
                    <default>{"true" if v == default else "false"}</default>
                    <label>{esc(v)}</label>
                </value>""" for v in values)
    return f"""        <valueSet>
            <restricted>{"true" if restricted else "false"}</restricted>
            <valueSetDefinition>
                <sorted>false</sorted>
{vs}
            </valueSetDefinition>
        </valueSet>"""

def field_xml(f):
    kw, t = f.kw, f.ftype
    p = [f"        <fullName>{f.api}</fullName>"]
    if t == "Checkbox":
        p.append(f"        <defaultValue>{'true' if kw.get('default') else 'false'}</defaultValue>")
    if t in ("Currency", "Number", "Percent") and "default" in kw:
        p.append(f"        <defaultValue>{kw['default']}</defaultValue>")
    if f.desc: p.append(f"        <description>{esc(f.desc)}</description>")
    p.append(f"        <externalId>{'true' if kw.get('externalId') else 'false'}</externalId>")
    p.append(f"        <label>{esc(f.label)}</label>")
    req = "true" if kw.get("required") else "false"
    if t == "Text":
        p += [f"        <length>{kw['length']}</length>", f"        <required>{req}</required>",
              "        <trackTrending>false</trackTrending>", "        <type>Text</type>",
              f"        <unique>{'true' if kw.get('unique') else 'false'}</unique>"]
    elif t in ("Phone", "Email", "Date", "DateTime", "Url"):
        p += [f"        <required>{req}</required>", "        <trackTrending>false</trackTrending>", f"        <type>{t}</type>"]
        if t == "Email": p.append("        <unique>false</unique>")
    elif t == "Checkbox":
        p += ["        <trackTrending>false</trackTrending>", "        <type>Checkbox</type>"]
    elif t in ("Currency", "Number", "Percent"):
        p += [f"        <precision>{kw.get('precision', 18)}</precision>", f"        <required>{req}</required>",
              f"        <scale>{kw.get('scale', 2)}</scale>", "        <trackTrending>false</trackTrending>", f"        <type>{t}</type>"]
        if t == "Number": p.append("        <unique>false</unique>")
    elif t == "Picklist":
        p += [f"        <required>{req}</required>", "        <trackTrending>false</trackTrending>", "        <type>Picklist</type>",
              picklist_xml(kw["values"], kw.get("default"), kw.get("restricted", True))]
    elif t == "LongTextArea":
        p += [f"        <length>{kw.get('length', 32768)}</length>", "        <trackTrending>false</trackTrending>",
              "        <type>LongTextArea</type>", f"        <visibleLines>{kw.get('visibleLines', 5)}</visibleLines>"]
    elif t == "Lookup":
        p += ["        <deleteConstraint>SetNull</deleteConstraint>", f"        <referenceTo>{kw['refTo']}</referenceTo>",
              f"        <relationshipLabel>{esc(kw.get('relLabel', f.label + 's'))}</relationshipLabel>",
              f"        <relationshipName>{kw['relName']}</relationshipName>", f"        <required>{req}</required>",
              "        <trackTrending>false</trackTrending>", "        <type>Lookup</type>"]
    elif t == "Formula":
        rtype = kw.get("rtype", "Currency")
        p.append(f"        <formula>{esc(kw['formula'])}</formula>")
        if rtype != "Text": p.append("        <formulaTreatBlanksAs>BlankAsZero</formulaTreatBlanksAs>")
        if rtype in ("Currency", "Number", "Percent"):
            p.append(f"        <precision>{kw.get('precision', 18)}</precision>")
        p.append("        <required>false</required>")
        if rtype in ("Currency", "Number", "Percent"):
            p.append(f"        <scale>{kw.get('scale', 2)}</scale>")
        p += ["        <trackTrending>false</trackTrending>", f"        <type>{rtype}</type>"]
    else:
        raise ValueError(t)
    return "    <fields>\n" + "\n".join(p) + "\n    </fields>"

def obj_xml(o):
    fields = "\n".join(field_xml(f) for f in o["fields"])
    if o["name_type"] == "AutoNumber":
        nf = f"""    <nameField>
        <displayFormat>{o['name_fmt']}</displayFormat>
        <label>{esc(o['name_label'])}</label>
        <type>AutoNumber</type>
    </nameField>"""
    else:
        nf = f"""    <nameField>
        <label>{esc(o['name_label'])}</label>
        <type>Text</type>
    </nameField>"""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<!--
{o['header']}
-->
<CustomObject xmlns="http://soap.sforce.com/2006/04/metadata">
    <deploymentStatus>Deployed</deploymentStatus>
    <enableActivities>false</enableActivities>
    <enableBulkApi>true</enableBulkApi>
    <enableFeeds>false</enableFeeds>
    <enableHistory>false</enableHistory>
    <enableReports>true</enableReports>
    <enableSearch>true</enableSearch>
    <enableSharing>true</enableSharing>
    <enableStreamingApi>true</enableStreamingApi>
{fields}
    <label>{esc(o['label'])}</label>
{nf}
    <pluralLabel>{esc(o['plural'])}</pluralLabel>
    <searchLayouts/>
    <sharingModel>Private</sharingModel>
    <externalSharingModel>Private</externalSharingModel>
</CustomObject>
"""

def lk(api, label, section, refTo, relName, desc="", relLabel=None, required=False, cache=True, sys=False):
    return F(api, label, section, "Lookup", desc, cache=cache, sys=sys, refTo=refTo, relName=relName,
             relLabel=relLabel or label + "s", required=required)

def tenant(relName, relLabel):
    return lk("Client__c", "Client", "Identity", TENANT, relName, "Tenant isolation anchor (D-034/D-035). Every Lambda query filters on this.", relLabel, required=True)

SNAP = lambda sec: [
    F("Customer_Name_at_Creation__c", "Customer Name at Creation", sec, "Text", "Snapshot at record creation (standard snapshot pattern).", sys=True, length=120),
    F("Address_at_Creation__c", "Address at Creation", sec, "Text", "Snapshot at record creation.", sys=True, length=255),
    F("Primary_Phone_at_Creation__c", "Primary Phone at Creation", sec, "Phone", "Snapshot at record creation.", sys=True),
    F("Primary_Email_at_Creation__c", "Primary Email at Creation", sec, "Email", "Snapshot at record creation.", sys=True),
]

OBJECTS = []

# ============================ Sundial_Estimate__c ============================
OBJECTS.append(dict(
    api="Sundial_Estimate__c", label="Service Estimate", plural="Service Estimates", key="estimate",
    table="sundial_estimate_cache", name_type="AutoNumber", name_fmt="EST-{00000}", name_label="Estimate Number",
    parent=("Sundial_Customer__c", "sundial_customer_sf_id"),
    search=(["name", "customer_name_at_creation", "template_name"], ["Name", "Customer_Name_at_Creation__c", "Template_Name__c"]),
    delete=False,
    header="""  Sundial_Estimate__c - the quote AND the living bill of work (D-072, docs/service-data-model.md).
  Can exist without a job (proposals); every job has exactly one (Service_Job__c is the 1:1 pair,
  set at conversion). Service Lines hang off THIS object, never the job. Money totals are
  computed by the estimate Lambda on every line write and stored here (D-072.8); the job
  reads them through cross-object formulas. Versions = append-only Version_Log__c JSON + a
  PDF per send (no Version object). Templates are estimates with Is_Template__c and no customer.""",
    fields=[
        lk("Sundial_Customer__c", "Sundial Customer", "Identity", "Sundial_Customer__c", "Service_Estimates", "Required unless Is_Template__c (templates have no customer; enforced in code - a required lookup would block templates).", "Service Estimates"),
        lk("Service_Job__c", "Service Job", "Identity", "Sundial_Service_Job__c", "Estimates", "The 1:1 pair. Null while the estimate is a proposal with no job; set by Create Job / quick-create. One job per estimate is Lambda-enforced.", "Estimates"),
        tenant("Sundial_Estimates", "Service Estimates"),
        *SNAP("Identity"),
        lk("Originating_Solar_Project__c", "Originating Solar Project", "Context", "Sundial_Solar__c", "Service_Estimates", "System being quoted on (no Asset object, D-065.1).", "Service Estimates"),
        lk("Originating_Roofing_Project__c", "Originating Roofing Project", "Context", "Sundial_Roofing__c", "Service_Estimates", "", "Service Estimates"),
        lk("Originating_Commercial_Project__c", "Originating Commercial Project", "Context", "Sundial_Commercial__c", "Service_Estimates", "Phase 3 object - REMOVE from the package if Sundial_Commercial__c is absent (verify script flags it).", "Service Estimates"),
        lk("Sold_By__c", "Sold By", "Context", "Sundial_User__c", "Sold_Estimates", "Commission attribution (with Markup below). Blank = office/no commission.", "Sold Estimates"),
        F("Is_Template__c", "Is Template", "Template", "Checkbox", "True = a named line set (Paige's Fronius/removal-reinstall/EV templates). No customer, never sent; 'Add from template' clones its lines.", default=False),
        F("Template_Name__c", "Template Name", "Template", "Text", "Display name in the template picker.", length=120),
        F("Status__c", "Status", "Status & Versions", "Picklist", "Draft -> Sent -> (Viewed) -> Approved | Declined | Expired; Invoiced when the job's invoice issues; Template never leaves Template. Re-sending after approval returns to Sent (Approved_Version__c keeps the last approved).", values=["Draft", "Sent", "Viewed", "Approved", "Declined", "Expired", "Invoiced", "Template"], default="Draft"),
        F("Version__c", "Version", "Status & Versions", "Number", "Incremented on every Send. 0 = never sent.", sys=True, precision=5, scale=0, default="0"),
        F("Version_Log__c", "Version Log", "Status & Versions", "LongTextArea", "Append-only JSON, Lambda-written, one entry per send: {version, sentAt, sentBy, sentVia, total, lines:[{itemCode,itemVersion,desc,qty,unitPrice,kind}], pdfKey}. THIS is the version history - no Version object (D-072.5).", cache=False, sys=True, length=131072, visibleLines=8),
        F("Last_Sent_At__c", "Last Sent At", "Status & Versions", "DateTime", "", sys=True),
        F("Last_Sent_Via__c", "Last Sent Via", "Status & Versions", "Picklist", "", sys=True, values=["Email", "SMS", "Both", "Manual"]),
        F("Last_Viewed_At__c", "Last Viewed At", "Status & Versions", "DateTime", "Hosted page opened (token hit).", sys=True),
        F("Approved_At__c", "Approved At", "Approval", "DateTime", "", sys=True),
        F("Approved_Version__c", "Approved Version", "Approval", "Number", "Which sent version the customer approved.", sys=True, precision=5, scale=0),
        F("Approved_Amount__c", "Approved Amount", "Approval", "Currency", "Total of the approved version (frozen - Total__c keeps moving as work is added).", sys=True, precision=16, scale=2),
        F("Approval_Method__c", "Approval Method", "Approval", "Picklist", "Online = accept button on the hosted page; Verbal = office logged it; Deposit Paid = paying the deposit is acceptance.", values=["Online", "Verbal", "Signed", "Deposit Paid"]),
        F("Approved_By_Name__c", "Approved By Name", "Approval", "Text", "Typed name on the hosted page, or the office user for Verbal.", length=120),
        F("Declined_Reason__c", "Declined Reason", "Approval", "Text", "", length=255),
        F("Valid_Until__c", "Valid Until", "Approval", "Date", "Send date + tenant validity days (GET FROM HARMON: the number). Past this the status flips to Expired by the nightly job."),
        F("Reminder_1_Sent_At__c", "Reminder 1 Sent At", "Approval", "DateTime", "+3d reminder (tenant config).", sys=True),
        F("Reminder_2_Sent_At__c", "Reminder 2 Sent At", "Approval", "DateTime", "+5d reminder; after this the estimate hands off to the NSA drip (D-065 amendment 1).", sys=True),
        F("Public_Token__c", "Public Token", "Hosted Page", "Text", "Unguessable token for the customer-facing estimate page (accept + card authorization at the bottom - D-072.7). External ID so the page route resolves it in one query.", cache=False, sys=True, length=64, externalId=True),
        F("Public_Token_Expires_At__c", "Public Token Expires At", "Hosted Page", "DateTime", "", cache=False, sys=True),
        F("Labor_Subtotal__c", "Labor Subtotal", "Money", "Currency", "Sum of Labor-kind lines + labor portion of Product lines. Lambda-computed on every line write (D-072.8).", sys=True, precision=16, scale=2),
        F("Material_Subtotal__c", "Material Subtotal", "Money", "Currency", "Sum of Material-kind lines + material portion of Product lines.", sys=True, precision=16, scale=2),
        F("Fee_Subtotal__c", "Fee Subtotal", "Money", "Currency", "Sum of Fee-kind lines.", sys=True, precision=16, scale=2),
        F("Subtotal__c", "Subtotal", "Money", "Currency", "Labor + Material + Fee, before discount/markup/tax.", sys=True, precision=16, scale=2),
        F("Discount_Scope__c", "Discount Scope", "Money", "Picklist", "Which subtotal the discount applies to (labor-only / material-only / both - D-065 amendment 4).", values=["Labor", "Material", "Both"], default="Both"),
        F("Discount_Type__c", "Discount Type", "Money", "Picklist", "", values=["Percent", "Amount"], default="Percent"),
        F("Discount_Value__c", "Discount Value", "Money", "Number", "The percent (10 = 10%) or the dollar amount, per Discount_Type__c.", precision=16, scale=2),
        F("Discount_Amount__c", "Discount Amount", "Money", "Currency", "Computed dollars off.", sys=True, precision=16, scale=2),
        F("Discount_Source__c", "Discount Source", "Money", "Picklist", "Service Plan = applied automatically from the customer's active plan; Manual = office entered.", values=["Manual", "Service Plan"], default="Manual"),
        F("Markup_Type__c", "Markup Type", "Money", "Picklist", "Hidden from the customer - only the total shows. Tim's commission mechanism with Sold_By__c.", values=["Percent", "Amount"], default="Percent"),
        F("Markup_Value__c", "Markup Value", "Money", "Number", "", precision=16, scale=2),
        F("Markup_Amount__c", "Markup Amount", "Money", "Currency", "Computed dollars added (never printed as a line).", sys=True, precision=16, scale=2),
        F("Tax_Rate__c", "Tax Rate", "Money", "Percent", "Resolved from the service-address city via the per-tenant AZ table - never a fixed default.", sys=True, precision=18, scale=3),
        F("Tax_Jurisdiction__c", "Tax Jurisdiction", "Money", "Text", "City/rate label used, for the printed document.", sys=True, length=80),
        F("Tax_Amount__c", "Tax Amount", "Money", "Currency", "Tax over Taxable lines after discount/markup allocation.", sys=True, precision=16, scale=2),
        F("Total__c", "Total", "Money", "Currency", "Subtotal - Discount + Markup + Tax. The number the customer sees.", sys=True, precision=16, scale=2),
        F("Deposit_Required__c", "Deposit Required", "Deposit", "Checkbox", "When true the hosted page charges the deposit at acceptance (and saves the card).", default=False),
        F("Deposit_Type__c", "Deposit Type", "Deposit", "Picklist", "", values=["Flat", "Percent"], default="Percent"),
        F("Deposit_Value__c", "Deposit Value", "Deposit", "Number", "", precision=16, scale=2),
        F("Deposit_Amount__c", "Deposit Amount", "Deposit", "Currency", "Computed from Deposit_Type/Value against Total.", sys=True, precision=16, scale=2),
        F("Deposit_Paid_At__c", "Deposit Paid At", "Deposit", "DateTime", "Mirrors the Deposit-type Service Payment.", sys=True),
        F("Scope_Summary__c", "Scope Summary", "Text", "LongTextArea", "Customer-facing paragraph printed above the lines. AI-drafted from notes, office-edited.", cache=False),
        F("Internal_Notes__c", "Internal Notes", "Text", "LongTextArea", "Append-only stamped entries (D-065.8). Never printed.", cache=False, length=131072, visibleLines=8),
        F("Created_In_Field__c", "Created In Field", "Origin", "Checkbox", "True when a tech created/extended it from the PWA (field estimate, D-065.7).", default=False),
        lk("Created_By_Service_Call__c", "Created By Service Call", "Origin", "Sundial_Service_Call__c", "Field_Estimates", "The visit during which the tech created the estimate.", "Field Estimates"),
    ],
))

# ============================ Sundial_Service_Job__c ============================
OBJECTS.append(dict(
    api="Sundial_Service_Job__c", label="Service Job", plural="Service Jobs", key="job",
    table="sundial_service_job_cache", name_type="AutoNumber", name_fmt="SVC-{00000}", name_label="Job Number",
    parent=("Sundial_Customer__c", "sundial_customer_sf_id"),
    search=(["name", "customer_name_at_creation", "billing_reference"], ["Name", "Customer_Name_at_Creation__c", "Billing_Reference__c"]),
    delete=False,
    header="""  Sundial_Service_Job__c - the work (was Sundial_Service__c; renamed D-072). One record per
  customer issue: intake, triage, scheduling context, ONE payer (Bill-To lives here - a second
  payer at the same address is a second job), time roll-ups from Service Calls (Flow, D-065.3).
  Every job has exactly one Sundial_Estimate__c (required lookup) - lines and money live there;
  the Estimate_* formulas below read through. Invoice number = job number (D-065 amendment 1).
  Job Status has NO estimate states - those are Sundial_Estimate__c.Status__c (D-072.3).
  Run scripts/verify-service-schema.mjs BEFORE deploying.""",
    fields=[
        lk("Sundial_Customer__c", "Sundial Customer", "Identity", "Sundial_Customer__c", "Service_Jobs", "Required - every job belongs to a customer (D-065.1: no Asset; history from this link + originating-project lookups).", "Service Jobs", required=True),
        lk("Estimate__c", "Estimate", "Identity", "Sundial_Estimate__c", "Service_Jobs", "REQUIRED: the job's living estimate (D-072.2). Quick-create makes it in the same transaction; Create Job links an existing one.", "Service Jobs", required=True),
        tenant("Sundial_Service_Jobs", "Service Jobs"),
        *SNAP("Identity"),
        F("Intake_Channel__c", "Intake Channel", "Intake", "Picklist", "Which door the request came in.", values=["Phone", "Email", "Web Form", "Monitoring Alert", "Manufacturer Referral", "Leasing Company", "Online Booking", "Estimate Conversion"]),
        F("Intake_Date__c", "Intake Date", "Intake", "DateTime", "When the request arrived (not record creation)."),
        F("Source_Email_From__c", "Source Email From", "Intake", "Email", "Email-intake path: the original sender.", cache=False),
        F("Source_Email_Reference__c", "Source Email Reference", "Intake", "Text", "Sender-side case/ticket reference.", cache=False, length=120),
        F("Needs_Intake_Review__c", "Needs Intake Review", "Intake", "Checkbox", "Set by the AI email-intake worker; office confirms the customer match before triage. AI never auto-schedules.", default=False),
        lk("Assigned_To__c", "Assigned To", "Ownership", "Sundial_User__c", "Assigned_Service_Jobs", "Current owner; feeds My Queue and the assignment notification.", "Assigned Service Jobs"),
        F("Status__c", "Status", "Status", "Picklist", "Job pipeline (service-workflows.md 2.1 minus the estimate states). Transitions guarded in Lambda code, not validation rules.", values=["New", "Triaging", "Remote Investigation", "Ready to Schedule", "Scheduled", "In Progress", "Awaiting Parts", "Awaiting Office Review", "Ready to Bill", "Invoiced", "Paid", "Closed"], default="New"),
        F("Resolution__c", "Resolution", "Status", "Picklist", "Terminal disposition; required (in code) when Status = Closed.", values=["Completed", "Resolved Remotely - No Charge", "Estimate Declined", "Cancelled", "Duplicate", "Referred Out"]),
        F("Status_Changed_At__c", "Status Changed At", "Status", "DateTime", "Stamped on every transition; feeds sat-too-long alerts.", sys=True),
        F("Priority__c", "Priority", "Status", "Picklist", "", values=["Low", "Standard", "High", "Emergency"], default="Standard"),
        F("Service_Type__c", "Service Type", "Classification", "Picklist", "", values=["Warranty", "Paid Service", "Monitoring Follow-up", "Maintenance", "Upgrade", "Partner Work Order"]),
        F("System_Ownership__c", "System Ownership", "Classification", "Picklist", "Leased / third-party steers the Bill To default.", values=["Customer Owned", "Leased", "Third-Party Owned"]),
        F("Issue_Description__c", "Issue Description", "Narrative", "LongTextArea", "Customer-reported issue at intake (board tooltip reads the first line).", ),
        F("Initial_Remote_Diagnosis__c", "Initial Remote Diagnosis", "Narrative", "LongTextArea", "Remote-first findings before any truck roll.", cache=False),
        F("Office_Notes__c", "Office Notes", "Narrative", "LongTextArea", "Internal office log - append-only stamped entries (D-065.8). Separate from tech notes per the 9/9 meeting; tech notes live on Service Calls and are shown here read-only.", cache=False, length=131072, visibleLines=8),
        F("Customer_Summary__c", "Customer Summary", "Narrative", "LongTextArea", "The customer-facing paragraph for the receipt and the photo job report. AI-drafted from the calls' Work Notes, office-edited before send.", cache=False),
        lk("Originating_Solar_Project__c", "Originating Solar Project", "Context", "Sundial_Solar__c", "Service_Jobs", "Installed system: specs, install date, photos read from here (no Asset object).", "Service Jobs"),
        lk("Originating_Roofing_Project__c", "Originating Roofing Project", "Context", "Sundial_Roofing__c", "Service_Jobs", "", "Service Jobs"),
        lk("Originating_Commercial_Project__c", "Originating Commercial Project", "Context", "Sundial_Commercial__c", "Service_Jobs", "Phase 3 object - REMOVE from the package if Sundial_Commercial__c is absent.", "Service Jobs"),
        F("Bill_To_Type__c", "Bill To Type", "Billing (one payer)", "Picklist", "THE payer for this job (D-072.6). No customer default, no per-visit override. 'Internal Warranty' is the tenant-neutral value for on-us warranty work.", values=BILL_TO, default="Customer"),
        F("Bill_To_Name__c", "Bill To Name", "Billing (one payer)", "Text", "Payer name when not the customer (leasing partner, manufacturer). Partner list is tenant config.", length=255),
        F("Billing_Reference__c", "Billing Reference", "Billing (one payer)", "Text", "Partner PO / work-order number. External ID (indexed) - 'SunRun calls with THEIR number'; printed on the invoice.", length=100, externalId=True),
        F("Payment_Status__c", "Payment Status", "Billing (one payer)", "Picklist", "Derived from Service Payments vs invoice total (Lambda/Flow).", sys=True, values=["None", "Deposit Paid", "Partially Paid", "Paid", "Refunded"], default="None"),
        F("Customer_Card_on_File__c", "Customer Card on File", "Billing (one payer)", "Checkbox", "A SetupIntent completed for this customer (mirror of the customer's Stripe state at last check). Cards live in Stripe, never here.", sys=True, default=False),
        F("Estimate_Status__c", "Estimate Status", "Estimate (read-through)", "Formula", "Cross-object formula - the job list shows it as a column.", sys=True, rtype="Text", formula="TEXT(Estimate__r.Status__c)"),
        F("Estimate_Total__c", "Estimate Total", "Estimate (read-through)", "Formula", "Cross-object formula onto the living estimate - no sync, no drift (D-072.3).", sys=True, rtype="Currency", precision=18, scale=2, formula="Estimate__r.Total__c"),
        F("Estimate_Approved_Amount__c", "Estimate Approved Amount", "Estimate (read-through)", "Formula", "What the customer last approved.", sys=True, rtype="Currency", precision=18, scale=2, formula="Estimate__r.Approved_Amount__c"),
        F("Estimate_Deposit_Amount__c", "Estimate Deposit Amount", "Estimate (read-through)", "Formula", "", sys=True, rtype="Currency", precision=18, scale=2, formula="Estimate__r.Deposit_Amount__c"),
        F("Total_Call_Count__c", "Total Call Count", "Time roll-ups", "Number", "Roll-up via record-triggered Flow on Service Call change (D-065.3). Migration writes totals directly.", sys=True, precision=5, scale=0),
        F("Total_Time_Minutes__c", "Total Time Minutes", "Time roll-ups", "Number", "Sum of Service Call durations across all techs (roll-up Flow).", sys=True, precision=9, scale=0),
        F("First_Scheduled_Start__c", "First Scheduled Start", "Time roll-ups", "DateTime", "Earliest Service Call Scheduled_Start (roll-up Flow) - list sorting.", sys=True),
        F("Geocode_Lat__c", "Geocode Latitude", "Location", "Number", "Service-address geocode, best-effort at intake; feeds geofence + travel hints.", sys=True, precision=18, scale=15),
        F("Geocode_Lon__c", "Geocode Longitude", "Location", "Number", "", sys=True, precision=18, scale=15),
        F("Geocode_Status__c", "Geocode Status", "Location", "Picklist", "Failed/Manual rows fall back to no geofence rather than blocking clock-ins.", sys=True, values=["Pending", "Geocoded", "Failed", "Manual"], default="Pending"),
        F("Street_View_Image_Key__c", "Street View Image Key", "Location", "Text", "S3 key of the street-view still fetched once at geocode time (9/9 ask). Blank = none available.", cache=False, sys=True, length=255),
    ],
))

# ============================ Sundial_Service_Call__c ============================
OBJECTS.append(dict(
    api="Sundial_Service_Call__c", label="Service Call", plural="Service Calls", key="servicecall",
    table="sundial_service_call_cache", name_type="AutoNumber", name_fmt="SC-{00000}", name_label="Call Number",
    parent=("Sundial_Service_Job__c", "sundial_service_job_sf_id"),
    search=None, delete=False,
    header="""  Sundial_Service_Call__c - one tech x one scheduled appointment, with its own clock and GPS
  (was Sundial_Service_Visit__c; renamed D-072). Multi-tech jobs are parallel records
  (D-065.2); D-027's dual-purpose design carries install visit types too. Parent lookups are
  deliberately NOT required in metadata (validation rule / Lambda guard ships with the
  workflows build so migration can bulk-load). Time fields are Lambda-written; the PWA has no
  time-edit surface (D-065.10). No Bill-To here any more: one job = one payer (D-072.6).""",
    fields=[
        F("Visit_Type__c", "Visit Type", "Identity", "Picklist", "Drives which parent lookup must be populated and which PWA tab renders the call (D-027).", values=["Service", "Solar Install", "Roofing Install", "Commercial Install"], default="Service"),
        lk("Sundial_Service_Job__c", "Service Job", "Identity", "Sundial_Service_Job__c", "Service_Calls", "Parent job - required (via validation) when Visit Type = Service.", "Service Calls"),
        lk("Sundial_Solar__c", "Solar Project", "Identity", "Sundial_Solar__c", "Install_Visits", "Parent when Visit Type = Solar Install.", "Install Visits"),
        lk("Sundial_Roofing__c", "Roofing Project", "Identity", "Sundial_Roofing__c", "Install_Visits", "Parent when Visit Type = Roofing Install.", "Install Visits"),
        lk("Sundial_Commercial__c", "Commercial Project", "Identity", "Sundial_Commercial__c", "Install_Visits", "Parent when Visit Type = Commercial Install. REMOVE from the package if Sundial_Commercial__c is absent.", "Install Visits"),
        lk("Tech__c", "Tech", "Identity", "Sundial_User__c", "Service_Calls", "ONE tech per call. A multi-tech job is parallel records under the same job (D-065.2).", "Service Calls"),
        tenant("Sundial_Service_Calls", "Service Calls"),
        F("Visit_Sub_Type__c", "Visit Sub Type", "Identity", "Picklist", "", values=["On-Site", "Remote", "Office Work", "In-Field", "Travel", "Prep"]),
        F("Scheduled_Start__c", "Scheduled Start", "Schedule", "DateTime", "Set by the dispatch board (block drop / resize)."),
        F("Scheduled_End__c", "Scheduled End", "Schedule", "DateTime", "Default length = sum of the estimate's labor Estimated_Hours (price book) or the service-type default."),
        F("Status__c", "Status", "Schedule", "Picklist", "En Route = 'on my way' (texts the customer when notify is on; ends the previous call's clock and starts drive time on this one).", values=["Scheduled", "En Route", "In Progress", "Complete", "Cancelled", "No-Show"], default="Scheduled"),
        F("Cancel_Reason__c", "Cancel Reason", "Schedule", "Text", "Required (in code) when the board cancels a call.", cache=False, length=255),
        F("Actual_Start__c", "Actual Start", "Time", "DateTime", "FIRST clock-in (device tap-time, not sync-time).", sys=True),
        F("Actual_End__c", "Actual End", "Time", "DateTime", "LAST clock-out; re-clock-in reopens the call.", sys=True),
        F("Duration_Minutes__c", "Duration Minutes", "Time", "Number", "Sum of the intervals in Clock Intervals - NOT end minus start. Feeds payroll AND billing.", sys=True, precision=9, scale=0),
        F("Clock_Intervals__c", "Clock Intervals", "Time", "LongTextArea", "Append-only JSON interval log: [{in, out, in_gps, out_gps, source}]. Manager corrections APPEND with actor + reason (D-065.10).", cache=False, sys=True, length=131072, visibleLines=8),
        F("Clock_In_Latitude__c", "Clock In Latitude", "Location", "Number", "GPS at first clock-in (null when no fix - never blocks).", sys=True, precision=18, scale=15),
        F("Clock_In_Longitude__c", "Clock In Longitude", "Location", "Number", "", sys=True, precision=18, scale=15),
        F("Clock_Out_Latitude__c", "Clock Out Latitude", "Location", "Number", "", sys=True, precision=18, scale=15),
        F("Clock_Out_Longitude__c", "Clock Out Longitude", "Location", "Number", "", sys=True, precision=18, scale=15),
        F("Geofence_Verified__c", "Geofence Verified", "Location", "Checkbox", "All clock events within the per-tenant radius of the service address OR the shop. TAG, not blocker.", sys=True, default=False),
        F("Work_Notes__c", "Work Notes", "Notes", "LongTextArea", "Tech notes - append-only stamped entries (D-065.8). Customer-facing candidates: the job's Customer Summary is AI-drafted from these.", cache=False, length=131072, visibleLines=10),
        F("Private_Notes__c", "Private Notes", "Notes", "LongTextArea", "Internal-only entries, same format. Never merges anywhere customer-visible.", cache=False, length=131072, visibleLines=10),
        F("Install_Milestones_Completed__c", "Install Milestones Completed", "Install", "LongTextArea", "Install visit types only.", cache=False),
        F("Checklist_Template_Key__c", "Checklist Template Key", "Checklist", "Text", "Per-tenant config key assigned at scheduling (D-065.9).", length=80),
        F("Checklist_State__c", "Checklist State", "Checklist", "LongTextArea", "Item completion snapshot (JSON). Required items gate Complete.", cache=False, sys=True),
        F("Photos_Count__c", "Photos Count", "Files", "Number", "Count of photos at SUNDIAL/{jobId}/photos/{callId}/ (metadata-derived).", sys=True, precision=4, scale=0),
    ],
))

# ============================ Sundial_Price_Book_Item__c ============================
OBJECTS.append(dict(
    api="Sundial_Price_Book_Item__c", label="Price Book Item", plural="Price Book Items", key="pricebookitem",
    table="sundial_price_book_item_cache", name_type="Text", name_label="Item Name",
    parent=None,
    search=(["name", "item_code", "description"], ["Name", "Item_Code__c", "Description__c"]),
    delete=False,
    header="""  Sundial_Price_Book_Item__c - the tenant-scoped, VERSIONED price book (D-072.4). NOT the
  standard Pricebook2/Product2 (org-wide, cannot be tenant-isolated). Item_Code__c is the
  stable id shared by every version; exactly one version per code is active (Lambda-enforced).
  'Update' in the portal = clone with the same code, Version+1, old row Is_Active=false and
  Superseded_By set. NEVER deleted (the permission set grants no delete); in-place edit only
  while no Service Line references the version. Lines snapshot price/cost/description anyway.""",
    fields=[
        tenant("Sundial_Price_Book_Items", "Price Book Items"),
        F("Item_Code__c", "Item Code", "Identity", "Text", "Stable line-item ID shared across versions (e.g. SVC-CALL-STD). Office-assigned; migration fills from HCP. External ID for lookup - NOT unique (versions share it).", length=40, externalId=True),
        F("Version__c", "Version", "Identity", "Number", "1, 2, 3... per Item Code.", sys=True, precision=5, scale=0, default="1"),
        F("Is_Active__c", "Is Active", "Identity", "Checkbox", "Portal Price Book list = active only. One active per Item Code.", default=True),
        lk("Superseded_By__c", "Superseded By", "Identity", "Sundial_Price_Book_Item__c", "Prior_Versions", "Set on the old version when Update creates the new one.", "Prior Versions", sys=True),
        F("Kind__c", "Kind", "Classification", "Picklist", "Product = labor + material on one item (both price splits filled). Drives kind-scoped discounts and taxability.", values=["Labor", "Material", "Product", "Fee"], default="Labor"),
        F("Category__c", "Category", "Classification", "Picklist", "GET FROM HARMON: final list, seeded from the HCP export. Unrestricted so the office can add values via the portal admin.", restricted=False, values=["Service Call", "Inverter", "Panel", "Battery", "EV Charger", "Electrical", "Roofing", "Inspection", "Cleaning", "Service Plan", "Materials", "Other"]),
        F("Description__c", "Description", "Content", "LongTextArea", "Customer-facing text printed on estimates/invoices.", length=4000, visibleLines=4),
        F("Internal_Notes__c", "Internal Notes", "Content", "LongTextArea", "Never printed.", cache=False),
        F("Unit_of_Measure__c", "Unit of Measure", "Pricing", "Picklist", "", values=["Each", "Hour", "Foot", "Lot"], default="Each"),
        F("Default_Quantity__c", "Default Quantity", "Pricing", "Number", "e.g. 1.5 hours for the standard service call.", precision=10, scale=2, default="1"),
        F("Estimated_Hours__c", "Estimated Hours", "Pricing", "Number", "Labor/Product items: default appointment duration when the job is scheduled.", precision=6, scale=2),
        F("Labor_Cost__c", "Labor Cost", "Pricing", "Currency", "Internal cost (margin reporting).", precision=16, scale=2),
        F("Material_Cost__c", "Material Cost", "Pricing", "Currency", "Internal cost.", precision=16, scale=2),
        F("Labor_Price__c", "Labor Price", "Pricing", "Currency", "Sell-side labor portion. Labor items fill only this; Product items fill both.", precision=16, scale=2),
        F("Material_Price__c", "Material Price", "Pricing", "Currency", "Sell-side material portion.", precision=16, scale=2),
        F("Price__c", "Price", "Pricing", "Formula", "Labor Price + Material Price - the one number the grid shows.", sys=True, rtype="Currency", precision=18, scale=2, formula="BLANKVALUE(Labor_Price__c,0)+BLANKVALUE(Material_Price__c,0)"),
        F("Taxable__c", "Taxable", "Pricing", "Checkbox", "Materials default true; labor per tenant tax config. GET FROM HEATHER: how HCP taxed labor vs materials.", default=True),
        F("HCP_Id__c", "HCP Id", "Migration", "Text", "Housecall Pro price-book id (migration key).", cache=False, length=64, externalId=True),
    ],
))

# ============================ Sundial_Service_Line__c ============================
OBJECTS.append(dict(
    api="Sundial_Service_Line__c", label="Service Line", plural="Service Lines", key="serviceline",
    table="sundial_service_line_cache", name_type="AutoNumber", name_fmt="SL-{00000}", name_label="Line Number",
    parent=("Estimate__c", "estimate_sf_id"),
    search=None, delete=True,
    header="""  Sundial_Service_Line__c - the junction: estimate x price-book item (+ quantity, price/cost
  snapshots, ad-hoc lines). Re-parented to Sundial_Estimate__c in D-072 (lines NEVER hang off
  the job). References a specific item VERSION and snapshots what it said at add time. The only
  service object with delete (removing a line while editing is a real row delete; sent
  versions survive in the estimate's Version_Log__c).""",
    fields=[
        lk("Estimate__c", "Estimate", "Identity", "Sundial_Estimate__c", "Service_Lines", "Parent estimate (required).", "Service Lines", required=True),
        lk("Price_Book_Item__c", "Price Book Item", "Identity", "Sundial_Price_Book_Item__c", "Service_Lines", "The specific item VERSION this line was added from. Blank = ad-hoc line (the portal offers save-to-price-book).", "Service Lines"),
        tenant("Sundial_Service_Lines", "Service Lines"),
        F("Kind__c", "Kind", "Content", "Picklist", "Copied from the item; required for ad-hoc lines. Scopes discounts.", values=["Labor", "Material", "Product", "Fee"], default="Labor"),
        F("Description__c", "Description", "Content", "Text", "Customer-facing line text (snapshot of the item description, editable per line).", length=255),
        F("Quantity__c", "Quantity", "Content", "Number", "", precision=10, scale=2, default="1"),
        F("Unit_of_Measure__c", "Unit of Measure", "Content", "Picklist", "", values=["Each", "Hour", "Foot", "Lot"], default="Each"),
        F("Unit_Price__c", "Unit Price", "Money", "Currency", "Snapshot of the item Price at add time; editable per line (Paige: 'sell it at 350').", precision=16, scale=2),
        F("Unit_Labor_Price__c", "Unit Labor Price", "Money", "Currency", "Snapshot split - lets a labor-only discount act on Product lines.", sys=True, precision=16, scale=2),
        F("Unit_Material_Price__c", "Unit Material Price", "Money", "Currency", "Snapshot split.", sys=True, precision=16, scale=2),
        F("Unit_Labor_Cost__c", "Unit Labor Cost", "Money", "Currency", "Snapshot for margin reporting.", sys=True, precision=16, scale=2),
        F("Unit_Material_Cost__c", "Unit Material Cost", "Money", "Currency", "Snapshot for margin reporting.", sys=True, precision=16, scale=2),
        F("Price_Overridden__c", "Price Overridden", "Money", "Checkbox", "Lambda-set when Unit Price differs from the item's price at add time (or was edited later).", sys=True, default=False),
        F("Line_Total__c", "Line Total", "Money", "Formula", "Quantity x Unit Price. Live formula.", sys=True, rtype="Currency", precision=18, scale=2, formula="BLANKVALUE(Quantity__c,0)*BLANKVALUE(Unit_Price__c,0)"),
        F("Taxable__c", "Taxable", "Money", "Checkbox", "Snapshot from the item; feeds tax at totals time.", default=True),
        F("Stage__c", "Stage", "Lifecycle", "Picklist", "Proposed = added but not yet in an approved version (field estimates land here); Approved; Completed; Removed = struck after a send (kept for the printed history).", values=["Proposed", "Approved", "Completed", "Removed"], default="Proposed"),
        F("Sort_Order__c", "Sort Order", "Lifecycle", "Number", "Grid order.", precision=5, scale=0),
        F("Show_Unit_Price__c", "Show Unit Price", "Lifecycle", "Checkbox", "Per-line display toggle; default from tenant config.", default=True),
        F("Source__c", "Source", "Lifecycle", "Picklist", "", values=["Price Book", "Template", "Ad hoc", "Field", "Migration"], default="Price Book"),
        lk("Added_By_Service_Call__c", "Added By Service Call", "Lifecycle", "Sundial_Service_Call__c", "Added_Service_Lines", "Set for field-added lines: which visit the tech was on.", "Added Service Lines"),
    ],
))

# ============================ Sundial_Service_Invoice__c ============================
OBJECTS.append(dict(
    api="Sundial_Service_Invoice__c", label="Service Invoice", plural="Service Invoices", key="serviceinvoice",
    table="sundial_service_invoice_cache", name_type="Text", name_label="Invoice Number",
    parent=("Service_Job__c", "service_job_sf_id"),
    search=(["name", "billing_reference"], ["Name", "Billing_Reference__c"]), delete=False,
    header="""  Sundial_Service_Invoice__c - the billing record, ONE per job (D-072.6; reissue after void =
  '-2'). Created from the estimate's lines at Ready to Bill and FROZEN. Customers receive a
  receipt + job report; partners receive this document (eventually generated in Acumatica -
  D-065 amendment 3). Name is TEXT: the Lambda assigns the job number. No Stripe fields here -
  money events are Sundial_Service_Payment__c rows.""",
    fields=[
        lk("Service_Job__c", "Service Job", "Identity", "Sundial_Service_Job__c", "Service_Invoices", "Parent job (required).", "Service Invoices", required=True),
        tenant("Sundial_Service_Invoices", "Service Invoices"),
        F("Status__c", "Status", "Status", "Picklist", "", values=["Draft", "Issued", "Sent", "Partially Paid", "Paid", "Void"], default="Draft"),
        F("Bill_To_Type__c", "Bill To Type", "Bill To (frozen)", "Picklist", "Frozen copy of the job's Bill To at issue.", sys=True, values=BILL_TO),
        F("Bill_To_Name__c", "Bill To Name", "Bill To (frozen)", "Text", "", sys=True, length=255),
        F("Billing_Reference__c", "Billing Reference", "Bill To (frozen)", "Text", "Partner PO/WO number, printed and searchable.", sys=True, length=100, externalId=True),
        F("Subtotal__c", "Subtotal", "Money (frozen)", "Currency", "", sys=True, precision=16, scale=2),
        F("Discount_Amount__c", "Discount Amount", "Money (frozen)", "Currency", "", sys=True, precision=16, scale=2),
        F("Tax_Rate__c", "Tax Rate", "Money (frozen)", "Percent", "", sys=True, precision=18, scale=3),
        F("Tax_Amount__c", "Tax Amount", "Money (frozen)", "Currency", "", sys=True, precision=16, scale=2),
        F("Total__c", "Total", "Money (frozen)", "Currency", "", sys=True, precision=16, scale=2),
        F("Paid_Amount__c", "Paid Amount", "Money (frozen)", "Currency", "Sum of Succeeded Service Payments (Payment - Refund); roll-up Flow.", sys=True, precision=16, scale=2, default="0"),
        F("Balance__c", "Balance", "Money (frozen)", "Formula", "Total - Paid.", sys=True, rtype="Currency", precision=18, scale=2, formula="BLANKVALUE(Total__c,0)-BLANKVALUE(Paid_Amount__c,0)"),
        F("Issued_At__c", "Issued At", "Dates", "DateTime", "", sys=True),
        F("Sent_At__c", "Sent At", "Dates", "DateTime", "Stamped on send / partner download.", sys=True),
        F("Due_Date__c", "Due Date", "Dates", "Date", "Partner terms (tenant config); blank for card-on-file jobs."),
        F("Paid_At__c", "Paid At", "Dates", "DateTime", "When Balance reached zero.", sys=True),
        F("PDF_S3_Key__c", "PDF S3 Key", "Documents", "Text", "SUNDIAL/{jobId}/{invoiceNumber}.pdf - appears in XFiles Pro + Dropbox automatically.", cache=False, sys=True, length=255),
        F("Acumatica_Ref__c", "Acumatica Ref", "Acumatica", "Text", "AR reference once the invoice exists in Acumatica (bridge: hand-entered; later: pushed).", length=50, externalId=True),
        F("Acumatica_Entered_At__c", "Acumatica Entered At", "Acumatica", "DateTime", "Bridge-period stamp: Heather's weekly digest = invoices where this is null."),
        F("Voided_At__c", "Voided At", "Void", "DateTime", "", sys=True),
        F("Void_Reason__c", "Void Reason", "Void", "Text", "", cache=False, length=255),
    ],
))

# ============================ Sundial_Service_Payment__c ============================
OBJECTS.append(dict(
    api="Sundial_Service_Payment__c", label="Service Payment", plural="Service Payments", key="servicepayment",
    table="sundial_service_payment_cache", name_type="AutoNumber", name_fmt="PAY-{00000}", name_label="Payment Number",
    parent=("Service_Job__c", "service_job_sf_id"),
    search=(["name", "reference"], ["Name", "Reference__c"]), delete=False,
    header="""  Sundial_Service_Payment__c - one row per money event on a job (D-072.1): deposit, final
  payment, refund, adjustment; card via Stripe, or check / partner remittance recorded by the
  office. Stripe_Payment_Intent_Id__c is UNIQUE - the webhook worker's idempotency key, so a
  redelivered event can never double-record. Heather's weekly digest and Beth's reports read
  these rows.""",
    fields=[
        lk("Service_Job__c", "Service Job", "Identity", "Sundial_Service_Job__c", "Service_Payments", "Required.", "Service Payments", required=True),
        lk("Invoice__c", "Invoice", "Identity", "Sundial_Service_Invoice__c", "Service_Payments", "Optional - deposits pre-date the invoice; the Lambda back-fills it at issue.", "Service Payments"),
        tenant("Sundial_Service_Payments", "Service Payments"),
        F("Type__c", "Type", "Money", "Picklist", "", values=["Deposit", "Payment", "Refund", "Adjustment"], default="Payment"),
        F("Method__c", "Method", "Money", "Picklist", "", values=["Card", "Check", "ACH", "Partner Remittance", "Other"], default="Card"),
        F("Amount__c", "Amount", "Money", "Currency", "Positive for money in; Refund rows are also positive (Type says the direction).", precision=16, scale=2),
        F("Status__c", "Status", "Money", "Picklist", "", values=["Pending", "Succeeded", "Failed", "Refunded"], default="Pending"),
        F("Stripe_Payment_Intent_Id__c", "Stripe Payment Intent Id", "Stripe", "Text", "UNIQUE external ID - webhook idempotency key.", sys=True, length=100, externalId=True, unique=True),
        F("Stripe_Charge_Id__c", "Stripe Charge Id", "Stripe", "Text", "", cache=False, sys=True, length=100),
        F("Stripe_Refund_Id__c", "Stripe Refund Id", "Stripe", "Text", "", cache=False, sys=True, length=100),
        F("Failure_Reason__c", "Failure Reason", "Stripe", "Text", "Stripe decline text for the office.", cache=False, sys=True, length=255),
        F("Received_At__c", "Received At", "Bookkeeping", "DateTime", "When the money actually landed (webhook time, or the check date the office enters)."),
        F("Reference__c", "Reference", "Bookkeeping", "Text", "Check number / partner remittance id.", length=100),
        lk("Recorded_By__c", "Recorded By", "Bookkeeping", "Sundial_User__c", "Recorded_Service_Payments", "Office user for manual entries; blank for webhook rows.", "Recorded Service Payments"),
        F("Acumatica_Applied_At__c", "Acumatica Applied At", "Bookkeeping", "DateTime", "When Heather applied it against the AR invoice (bridge) / when the push did."),
        F("Notes__c", "Notes", "Bookkeeping", "Text", "", cache=False, length=255),
    ],
))

# ============================ Sundial_Customer__c (ONE field, not whole-object) ============================
CUSTOMER_FIELD = F("Stripe_Customer_Id__c", "Stripe Customer Id", "Payments", "Text",
                   "Stripe customer reference (card on file via SetupIntent; also the Service Club subscription customer). Cards are vaulted in Stripe, never in Salesforce. Deployed as a single CustomField, never a whole-object deploy of Sundial_Customer__c.",
                   length=100, externalId=True)

# ---------------------------------------------------------------------------
# Emit .object files
# ---------------------------------------------------------------------------
if NO_COMMERCIAL:
    for o in OBJECTS:
        o["fields"] = [f for f in o["fields"] if f.kw.get("refTo") != "Sundial_Commercial__c"]
    print("--no-commercial: omitted the three Sundial_Commercial__c lookups")
for o in OBJECTS:
    open(f"{PKG}/objects/{o['api']}.object", "w").write(obj_xml(o))

open(f"{PKG}/objects/Sundial_Customer__c.object", "w").write(f"""<?xml version="1.0" encoding="UTF-8"?>
<!--
  Sundial_Customer__c - ONE NEW FIELD ONLY (D-072). package.xml lists it as a CustomField
  member, so this file adds Stripe_Customer_Id__c and touches nothing else on the object.
  NEVER convert this to a whole-object deploy: it would overwrite the live object's settings.
-->
<CustomObject xmlns="http://soap.sforce.com/2006/04/metadata">
{field_xml(CUSTOMER_FIELD)}
</CustomObject>
""")

# ---------------------------------------------------------------------------
# package.xml
# ---------------------------------------------------------------------------
members = "\n".join(f"        <members>{o['api']}</members>" for o in OBJECTS)
open(f"{PKG}/package.xml", "w").write(f"""<?xml version="1.0" encoding="UTF-8"?>
<!--
  service-objects v2 - the seven Service Operations objects (Phase 2, D-072), one new field
  on Sundial_Customer__c, and the integration-user permission set.

  BEFORE DEPLOYING: run  node scripts/verify-service-schema.mjs
  - If any of the seven objects ALREADY EXISTS, stop and reconcile (whole-object deploy).
  - If Sundial_Commercial__c is absent, remove the three commercial lookups first
    (Originating_Commercial_Project__c on estimate + job, Sundial_Commercial__c on the call)
    AND their fieldPermissions entries - a lookup to a missing object fails the whole deploy.

  DEPLOY: zip this folder's CONTENTS (package.xml at zip root; Linux/WSL zip or Explorer
  Send-to, NEVER PowerShell 5.1 Compress-Archive) -> Workbench -> Migration -> Deploy ->
  Single Package -> CHECK ONLY first, expect 9/9 components (7 objects + 1 field + 1
  permission set), then deploy for real. Then assign the permission set, re-run verify.
-->
<Package xmlns="http://soap.sforce.com/2006/04/metadata">
    <types>
{members}
        <name>CustomObject</name>
    </types>
    <types>
        <members>Sundial_Customer__c.Stripe_Customer_Id__c</members>
        <name>CustomField</name>
    </types>
    <types>
        <members>Sundial_Service_Objects</members>
        <name>PermissionSet</name>
    </types>
    <version>62.0</version>
</Package>
""")

# ---------------------------------------------------------------------------
# Permission set
# ---------------------------------------------------------------------------
def fp(obj, api, editable):
    return f"""    <fieldPermissions>
        <editable>{'true' if editable else 'false'}</editable>
        <field>{obj}.{api}</field>
        <readable>true</readable>
    </fieldPermissions>"""

fps, skipped_required = [], []
for o in OBJECTS:
    for f in o["fields"]:
        if f.kw.get("required"):
            skipped_required.append(f"{o['api']}.{f.api}")
            continue
        fps.append(fp(o["api"], f.api, f.ftype != "Formula"))
fps.append(fp("Sundial_Customer__c", CUSTOMER_FIELD.api, True))
ops = "\n".join(f"""    <objectPermissions>
        <allowCreate>true</allowCreate>
        <allowDelete>{'true' if o['delete'] else 'false'}</allowDelete>
        <allowEdit>true</allowEdit>
        <allowRead>true</allowRead>
        <modifyAllRecords>false</modifyAllRecords>
        <object>{o['api']}</object>
        <viewAllRecords>false</viewAllRecords>
    </objectPermissions>""" for o in OBJECTS)
open(f"{PKG}/permissionsets/Sundial_Service_Objects.permissionset", "w").write(f"""<?xml version="1.0" encoding="UTF-8"?>
<!--
  Object + field permissions on everything service-objects creates, for the Sundial
  INTEGRATION USER (assign to it after deploy). Same pattern as v6-access-model.

  allowDelete is TRUE only on Sundial_Service_Line__c. Jobs close, calls cancel, invoices
  void, estimates decline/expire, payments refund, and PRICE BOOK ITEMS DEACTIVATE -
  none of those are deletes, so a code mistake fails instead of destroying history (D-072.4).

  REQUIRED fields carry no fieldPermissions entries ON PURPOSE (Salesforce forbids FLS on
  required fields; they are always readable/writable):
  {", ".join(skipped_required)}.
  Formula fields are readable-only.

  Portal USERS need nothing here - they never touch Salesforce directly (D-002).
-->
<PermissionSet xmlns="http://soap.sforce.com/2006/04/metadata">
    <label>Sundial Service Objects</label>
    <description>Phase 2 (D-072): object and field access for the seven Service Operations objects + Sundial_Customer__c.Stripe_Customer_Id__c, for the Sundial integration user.</description>
{ops}
{chr(10).join(fps)}
</PermissionSet>
""")

# ---------------------------------------------------------------------------
# SQL cache tables
# ---------------------------------------------------------------------------
def col(api): return api[:-3].lower() if api.endswith("__c") else api.lower()
def pgtype(f):
    t = f.ftype
    if t == "Lookup": return "text"
    if t in ("Currency", "Number", "Percent"): return "numeric"
    if t == "Formula": return "text" if f.kw.get("rtype") == "Text" else "numeric"
    if t == "Checkbox": return "boolean"
    if t == "DateTime": return "timestamptz"
    if t == "Date": return "date"
    return "text"

INDEXES = {
    "estimate": [("customer", "(client_sf_id, sundial_customer_sf_id)"), ("job", "(client_sf_id, service_job_sf_id)"), ("status", "(client_sf_id, status)"), ("template", "(client_sf_id, is_template) where is_template = true")],
    "job": [("customer", "(client_sf_id, sundial_customer_sf_id)"), ("status", "(client_sf_id, status)"), ("queue", "(client_sf_id, assigned_to_sf_id, status)"), ("billing_ref", "(client_sf_id, billing_reference)"), ("estimate", "(client_sf_id, estimate_sf_id)")],
    "servicecall": [("window", "(client_sf_id, scheduled_start)"), ("tech_window", "(client_sf_id, tech_sf_id, scheduled_start)"), ("job", "(client_sf_id, sundial_service_job_sf_id)")],
    "pricebookitem": [("active", "(client_sf_id, is_active) where is_active = true"), ("code", "(client_sf_id, item_code, version desc)")],
    "serviceline": [("estimate", "(client_sf_id, estimate_sf_id, sort_order)"), ("item", "(client_sf_id, price_book_item_sf_id)")],
    "serviceinvoice": [("job", "(client_sf_id, service_job_sf_id)"), ("status", "(client_sf_id, status)"), ("acumatica_pending", "(client_sf_id, issued_at) where acumatica_entered_at is null")],
    "servicepayment": [("job", "(client_sf_id, service_job_sf_id)"), ("invoice", "(client_sf_id, invoice_sf_id)"), ("received", "(client_sf_id, received_at desc)")],
}

for o in OBJECTS:
    t, key = o["table"], o["key"]
    rows = []
    for f in o["fields"]:
        if not f.cache or f.api == "Client__c": continue
        c = col(f.api) + ("_sf_id" if f.ftype == "Lookup" else "")
        rows.append((c, pgtype(f), f"{f.api}" + (f" ({f.desc[:60]}...)" if len(f.desc) > 60 else (f" ({f.desc})" if f.desc else ""))))
    width = max(len(r[0]) for r in rows) + 2
    body = "\n".join(f"  {r[0]:<{width}} {r[1] + ',':<12} -- {r[2]}" for r in rows)
    name_comment = "Name (auto-number)" if o["name_type"] == "AutoNumber" else "Name (text)"
    idx = "\n".join(f"create index if not exists idx_{t}_{n}\n  on {t} {expr};" for n, expr in INDEXES[key])
    open(f"{ROOT}/sql/{t}.sql", "w").write(f"""-- {t} — Supabase cache table for {o['api']} ({o['label']}, D-072).
--
-- Mirrors the sundial_roofing_cache pattern: the display subset the portal's list/board/
-- grid surfaces render, plus the standard control columns every Sundial cache table
-- carries. Detail views use GET /sf/{key}/{{id}}?full=true (describe-driven, live).
--
-- HOW THIS POPULATES: column NAMES match sfFieldToColumn() (strip __c, lowercase;
-- reference fields get _sf_id). The read and sync Lambdas select ONLY Salesforce fields
-- whose mapped column exists here, so creating this table makes {key} records populate via
-- read-through and the scheduled sundial-cache-sync job. Add a column later and it starts
-- caching. GENERATED from the same spec as the .object files - edit the spec, not this.
--
-- TENANT ISOLATION: client_sf_id is the isolation key (D-035), NOT NULL;
-- tenant_id is the human slug (label only).

create table if not exists {t} (
  -- Control / identity
  {'sf_id':<{width}} text primary key,          -- Salesforce record Id
  {'client_sf_id':<{width}} text not null,             -- Client__c (tenant isolation key)
  {'tenant_id':<{width}} text,                      -- Client__r.Name slug (label only)

  -- Display subset
  {'name':<{width}} text,        -- {name_comment}
{body}

  {'created_date':<{width}} timestamptz, -- CreatedDate (list ordering)
  {'last_synced_at':<{width}} timestamptz not null default now(),
  {'cache_version':<{width}} integer not null default 1,
  {'is_stale':<{width}} boolean not null default false
);

create index if not exists idx_{t}_tenant
  on {t} (client_sf_id);
{idx}
create index if not exists idx_{t}_recent
  on {t} (client_sf_id, created_date desc nulls last, sf_id);
create index if not exists idx_{t}_stale
  on {t} (client_sf_id, is_stale) where is_stale = true;
""")

# ---------------------------------------------------------------------------
# Workbooks (generator format: Section / Field Label / API Name / Type / Description)
# ---------------------------------------------------------------------------
NAVY, BLUE, YELLOW = "1F3864", "D9E1F2", "FFF2CC"
def typelabel(f):
    t, kw = f.ftype, f.kw
    if t == "Text": return f"Text({kw['length']})" + (" · External ID" if kw.get("externalId") else "") + (" · Unique" if kw.get("unique") else "")
    if t == "Lookup": return f"Lookup({kw['refTo']})" + (" · Required" if kw.get("required") else "")
    if t == "Picklist": return "Picklist: " + " / ".join(kw["values"]) + (f" (default {kw['default']})" if kw.get("default") else "")
    if t in ("Currency", "Number", "Percent"): return f"{t}({kw.get('precision',18)},{kw.get('scale',2)})"
    if t == "Formula": return f"Formula ({kw.get('rtype','Currency')})"
    if t == "LongTextArea": return f"Long Text({kw.get('length',32768)})"
    if t == "Checkbox": return f"Checkbox (default {'true' if kw.get('default') else 'false'})"
    return t

def sheet(wb, title, o):
    ws = wb.create_sheet(title[:31])
    ws.append(["Section", "Field Label", "API Name", "Type", "Description / Calculation"])
    for c in ws[1]:
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    nm = "Auto Number" if o["name_type"] == "AutoNumber" else "Text(80)"
    ws.append(["Identity", o["name_label"], "Name", nm, o.get("name_fmt", "Lambda-assigned: job number, then -2 on reissue.") if o["name_type"] == "AutoNumber" else "Lambda-assigned."])
    for f in o["fields"]:
        ws.append([f.section, f.label, f.api, typelabel(f), f.desc + ("" if f.cache else " [not cached — detail only]")])
        if f.sys:
            for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=YELLOW)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate([26, 30, 34, 44, 66], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws

def legend(wb):
    ws = wb.create_sheet("Legend")
    ws.append(["Legend", ""])
    ws.append(["Yellow rows", "System-maintained: written by a Lambda, Flow, or webhook — never edited by a portal user directly. The generator renders them read-only."])
    ws.append(["[not cached]", "Field exists in Salesforce but has no column in the Supabase cache table; detail views read it live (?full=true)."])
    ws.append(["Source", "docs/service-data-model.md (D-072). Generated with the .object/.sql files from one spec — regenerate rather than hand-edit."])
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
    ws["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100

by_api = {o["api"]: o for o in OBJECTS}
books = {
    "Sundial_Estimate_Fields_by_Section.xlsx": [("Estimate Fields", "Sundial_Estimate__c"), ("Service Line Fields", "Sundial_Service_Line__c"), ("Price Book Item Fields", "Sundial_Price_Book_Item__c")],
    "Sundial_Service_Job_Fields_by_Section.xlsx": [("Service Job Fields", "Sundial_Service_Job__c"), ("Service Invoice Fields", "Sundial_Service_Invoice__c"), ("Service Payment Fields", "Sundial_Service_Payment__c")],
    "Sundial_Service_Call_Fields_by_Section.xlsx": [("Service Call Fields", "Sundial_Service_Call__c")],
}
for fn, sheets in books.items():
    wb = Workbook(); wb.remove(wb.active)
    for title, api in sheets: sheet(wb, title, by_api[api])
    legend(wb)
    wb.save(f"{ROOT}/docs/{fn}")

# ---------------------------------------------------------------------------
# spec.json for the registry edits + counts
# ---------------------------------------------------------------------------
spec = {o["key"]: dict(sfObject=o["api"], cacheTable=o["table"], parent=o["parent"], search=o["search"], delete=o["delete"],
                       fields=len(o["fields"]), cached=sum(1 for f in o["fields"] if f.cache)) for o in OBJECTS}
json.dump(spec, open(f"{PKG}/spec.json", "w"), indent=2)
total = sum(len(o["fields"]) for o in OBJECTS)
print(f"objects={len(OBJECTS)} fields={total} (+1 customer field) fls_entries={len(fps)} required_skipped={len(skipped_required)}")
for k, v in spec.items(): print(f"  {k:15} {v['sfObject']:32} fields={v['fields']:3} cached={v['cached']:3} table={v['cacheTable']}")
